import os
import re
import json
import time
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
from pathlib import Path
from dotenv import load_dotenv
from openai import AzureOpenAI
from qdrant_client import QdrantClient
from qdrant_client.models import Filter, FieldCondition, MatchValue

load_dotenv()

# ── Config ────────────────────────────────────────────────────
AZURE_API_KEY     = os.getenv("AZURE_OPENAI_API_KEY")
AZURE_ENDPOINT    = os.getenv("AZURE_OPENAI_ENDPOINT")
AZURE_API_VERSION = os.getenv("AZURE_OPENAI_API_VERSION")
EMBEDDING_MODEL   = "text-embedding-3-large"
CHAT_MODEL        = "gpt-4.1"
COLLECTION_NAME   = "hotel_contracts"
QDRANT_PATH       = "./qdrant_local"
OUTPUT_DIR        = Path("./reports")
TOP_K             = 5     # retrieve top 5 chunks per field

OUTPUT_DIR.mkdir(exist_ok=True)

openai_client = AzureOpenAI(
    api_key        = AZURE_API_KEY,
    api_version    = AZURE_API_VERSION,
    azure_endpoint = AZURE_ENDPOINT
)
qdrant_client = QdrantClient(path=QDRANT_PATH)


# ─────────────────────────────────────────────────────────────
# FIELD DEFINITIONS
# ─────────────────────────────────────────────────────────────
# Each field has:
#   retrieval_query  → what to search in vector DB
#   extract_prompt   → what to tell LLM to extract

FIELDS = [
    {
        "name":            "hotel_name",
        "label":           "Hotel Name",
        "retrieval_query": "name of hotel property",
        "extract_prompt":  "Extract the exact name of the hotel. Return only the name, nothing else."
    },
    {
        "name":            "type_of_agreement",
        "label":           "Agreement Type",
        "retrieval_query": "type of agreement hotel operating lease license agreement snapshot",
        "extract_prompt":  "What type of agreement is this? Return one of: 'Hotel Operating Agreement', 'Lease Agreement', 'License Agreement', or 'Management Agreement'. If not clear, return 'Not specified'."
    },
    {
        "name":            "hotel_opening_date",
        "label":           "Opening Date",
        "retrieval_query": "hotel opening date commencement operations started",
        "extract_prompt":  "Extract the exact hotel opening date. Return only the date (e.g., 'March 7, 2022'). If not specified, return 'N/A'."
    },
    {
        "name":            "no_of_rooms",
        "label":           "No. of Rooms",
        "retrieval_query": "number of rooms guest rooms suites total rooms",
        "extract_prompt":  "Extract the exact number of rooms. Return only the number (e.g., '325'). If not specified, return 'N/A'."
    },
    {
        "name":            "no_of_restaurants",
        "label":           "No. of Restaurants",
        "retrieval_query": "number of restaurants food beverage outlets dining",
        "extract_prompt":  "Extract the number of restaurants or food & beverage outlets. Return only the number. If not mentioned anywhere in the data, return 'Not specified'."
    },
    {
        "name":            "no_of_spas",
        "label":           "No. of SPAs",
        "retrieval_query": "number of spa wellness health club facilities",
        "extract_prompt":  "Extract the number of spas or wellness facilities. Return only the number. If not mentioned anywhere in the data, return 'Not specified'."
    },
    {
        "name":            "management_fee",
        "label":           "Management Fee",
        "retrieval_query": "management fee base management fee percentage revenue",
        "extract_prompt":  "Extract the management fee percentage AND its basis. Example: '1.0% of Total Revenue'. Include the full description. If N/A, return 'N/A'."
    },
    {
        "name":            "incentive_fee",
        "label":           "Incentive Fee",
        "retrieval_query": "incentive fee performance fee GOP AGOP percentage",
        "extract_prompt":  "Extract the incentive fee structure. Include percentages and conditions. Example: '4% of AGOP if GOP <= 40%, 7% if 40-50%'. If N/A, return 'N/A'."
    },
    {
        "name":            "fee_threshold",
        "label":           "Fee Threshold",
        "retrieval_query": "fee threshold cap limit maximum fee combined",
        "extract_prompt":  "Extract the fee threshold or cap. Include the percentage or amount. If N/A, return 'N/A'."
    },
    {
        "name":            "sales_marketing_fee",
        "label":           "Sales & Marketing Fee",
        "retrieval_query": "sales marketing fee percentage room revenue",
        "extract_prompt":  "Extract the sales and marketing fee. Include percentage and basis. Example: '1.0% of Room Revenue'. If N/A, return 'N/A'."
    },
]


# ─────────────────────────────────────────────────────────────
# STEP 1 — GET ALL CONTRACTS
# ─────────────────────────────────────────────────────────────

def get_all_contract_ids() -> list[str]:
    """Get unique contract IDs from Qdrant."""
    all_points = []
    offset = None

    while True:
        results = qdrant_client.scroll(
            collection_name = COLLECTION_NAME,
            limit           = 100,
            offset          = offset,
            with_payload    = True,
            with_vectors    = False
        )
        points, next_offset = results
        all_points.extend(points)
        if next_offset is None:
            break
        offset = next_offset

    contract_ids = list(set(
        p.payload.get("contract_id", "") for p in all_points
    ))
    return sorted(contract_ids)


# ─────────────────────────────────────────────────────────────
# STEP 2 — RETRIEVE RELEVANT CHUNKS PER FIELD
# ─────────────────────────────────────────────────────────────

def embed_query(query: str) -> list[float]:
    """Embed a retrieval query."""
    response = openai_client.embeddings.create(
        input = [query],
        model = EMBEDDING_MODEL
    )
    return response.data[0].embedding


def retrieve_chunks(contract_id: str, query: str, top_k: int = TOP_K) -> list[dict]:
    """
    Semantic search WITHIN a specific contract.
    Returns top_k most relevant chunks for the query.
    """
    query_vector = embed_query(query)

    results = qdrant_client.query_points(
        collection_name = COLLECTION_NAME,
        query           = query_vector,
        query_filter    = Filter(
            must=[FieldCondition(
                key   = "contract_id",
                match = MatchValue(value=contract_id)
            )]
        ),
        limit        = top_k,
        with_payload = True
    ).points

    return [
        {
            "section": r.payload.get("section", ""),
            "key":     r.payload.get("key", ""),
            "value":   r.payload.get("value", ""),
            "score":   round(r.score, 4)
        }
        for r in results
    ]


# ─────────────────────────────────────────────────────────────
# STEP 3 — LLM EXTRACTION FROM RETRIEVED CHUNKS
# ─────────────────────────────────────────────────────────────

EXTRACT_SYSTEM = """
You are a precise data extractor for hotel contracts.

Rules:
- Extract ONLY from the provided data. NEVER fabricate or guess.
- If the information is not in the provided data, say "Not specified".
- If the value is N/A in the data, return "N/A".
- Be concise. Return only the extracted value.
- Do not add explanations or commentary.
"""


def extract_field(field: dict, chunks: list[dict]) -> str:
    """
    Send retrieved chunks to LLM, extract one specific field.
    """
    if not chunks:
        return "Not specified"

    # build context from retrieved chunks
    context_lines = []
    for c in chunks:
        context_lines.append(
            f"[Section: {c['section']}] [Key: {c['key']}] "
            f"[Relevance: {c['score']}]\n{c['value']}"
        )
    context = "\n\n".join(context_lines)

    user_prompt = f"""
Retrieved contract data (most relevant chunks):
---
{context}
---

Task: {field['extract_prompt']}
"""

    try:
        response = openai_client.chat.completions.create(
            model       = CHAT_MODEL,
            messages    = [
                {"role": "system", "content": EXTRACT_SYSTEM},
                {"role": "user",   "content": user_prompt}
            ],
            temperature = 0,
            max_tokens  = 300
        )
        return response.choices[0].message.content.strip()

    except Exception as e:
        print(f"    ❌ LLM error: {e}")
        return "ERROR"


# ─────────────────────────────────────────────────────────────
# STEP 4 — PROCESS SINGLE CONTRACT
# ─────────────────────────────────────────────────────────────

def process_contract(contract_id: str) -> dict:
    """
    For one contract:
    - For each field, retrieve relevant chunks
    - Extract value using LLM
    - Return structured dict
    """
    record = {"contract_id": contract_id}

    for field in FIELDS:
        # RETRIEVE — semantic search within this contract
        chunks = retrieve_chunks(contract_id, field["retrieval_query"])

        # EXTRACT — LLM reads only relevant chunks
        value = extract_field(field, chunks)
        record[field["name"]] = value

        # show progress
        print(f"    {field['label']:20s} → {value[:60]}")

    return record


# ─────────────────────────────────────────────────────────────
# STEP 5 — EXPORT TO EXCEL
# ─────────────────────────────────────────────────────────────

EXCEL_COLUMNS = [
    ("#",                   5),
    ("Hotel Name",         30),
    ("Agreement Type",     25),
    ("Opening Date",       18),
    ("No. of Rooms",       15),
    ("No. of Restaurants", 18),
    ("No. of SPAs",        13),
    ("Management Fee",     30),
    ("Incentive Fee",      40),
    ("Fee Threshold",      35),
    ("Sales & Marketing",  30),
    ("Contract ID",        50),
]


def export_report(all_data: list[dict]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hotel Contracts Report"

    header_font  = Font(bold=True, color="FFFFFF", size=10)
    header_fill  = PatternFill("solid", fgColor="1F4E79")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border  = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0')
    )

    # ── title row ─────────────────────────────────────────────
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(EXCEL_COLUMNS))
    title = ws.cell(row=1, column=1,
        value=f"Hotel Contracts Report — {datetime.now().strftime('%d %b %Y %H:%M')}")
    title.font      = Font(bold=True, size=14, color="1F4E79")
    title.fill      = PatternFill("solid", fgColor="D6E4F0")
    title.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    # ── headers ───────────────────────────────────────────────
    for col, (header, width) in enumerate(EXCEL_COLUMNS, 1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = header_align
        cell.border    = thin_border
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    ws.row_dimensions[2].height = 22

    # ── data rows ─────────────────────────────────────────────
    alt_fill = PatternFill("solid", fgColor="EBF3FB")
    field_names = [f["name"] for f in FIELDS]

    for i, record in enumerate(all_data, 1):
        row  = i + 2
        fill = alt_fill if i % 2 == 0 else None

        # row number
        cell = ws.cell(row=row, column=1, value=i)
        cell.alignment = Alignment(horizontal="center", vertical="top")
        cell.border = thin_border
        if fill: cell.fill = fill

        # field values
        for col_offset, fname in enumerate(field_names, 2):
            cell = ws.cell(row=row, column=col_offset, value=record.get(fname, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border
            if fill: cell.fill = fill

        # contract id (last column)
        cell = ws.cell(row=row, column=len(EXCEL_COLUMNS), value=record.get("contract_id", ""))
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = thin_border
        if fill: cell.fill = fill

    # ── freeze + filter ───────────────────────────────────────
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{openpyxl.utils.get_column_letter(len(EXCEL_COLUMNS))}{len(all_data)+2}"

    # ── save ──────────────────────────────────────────────────
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filepath  = OUTPUT_DIR / f"hotel_contracts_report_{timestamp}.xlsx"
    wb.save(filepath)
    return filepath


# ─────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────

def generate_report():
    print(f"\n{'═'*60}")
    print(f"  HOTEL CONTRACTS REPORT GENERATOR")
    print(f"  Approach: Retrieve per field → LLM extract")
    print(f"{'═'*60}")

    contract_ids = get_all_contract_ids()
    print(f"\nFound {len(contract_ids)} contracts\n")

    all_data = []

    for i, cid in enumerate(contract_ids, 1):
        print(f"\n[{i}/{len(contract_ids)}] {cid[:55]}...")
        print(f"  {'─'*50}")

        record = process_contract(cid)
        all_data.append(record)
        time.sleep(0.3)

    filepath = export_report(all_data)

    print(f"\n{'═'*60}")
    print(f"✅ Report: {filepath}")
    print(f"   Contracts: {len(all_data)}")
    print(f"   Fields: {len(FIELDS)} per contract")
    print(f"   Total LLM calls: {len(all_data) * len(FIELDS)}")
    print(f"{'═'*60}")


if __name__ == "__main__":
    try:
        generate_report()
    finally:
        qdrant_client.close()